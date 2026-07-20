"""Extract plain text from uploaded documents for the Grafix prompt."""

from __future__ import annotations

import io
import os
import re
import subprocess
import tempfile
from pathlib import Path
from typing import Callable

ALLOWED_EXTENSIONS = {".txt", ".docx", ".doc", ".rtf", ".pdf"}
MAX_UPLOAD_BYTES = int(os.environ.get("GRAFIX_MAX_UPLOAD_MB", "12")) * 1024 * 1024


class DocumentExtractError(ValueError):
    """Raised when a document cannot be parsed into text."""


def normalize_ext(filename: str) -> str:
    return Path(filename or "").suffix.lower()


def extract_text(filename: str, data: bytes) -> str:
    if not data:
        raise DocumentExtractError("Пустой файл")
    if len(data) > MAX_UPLOAD_BYTES:
        raise DocumentExtractError(
            f"Файл слишком большой (лимит {MAX_UPLOAD_BYTES // (1024 * 1024)} МБ)"
        )

    ext = normalize_ext(filename)
    if ext not in ALLOWED_EXTENSIONS:
        raise DocumentExtractError(
            "Поддерживаются только .txt, .docx, .doc, .rtf, .pdf"
        )

    # Misnamed / renamed Office Open XML
    if ext == ".doc" and data[:2] == b"PK":
        ext = ".docx"

    extractors: dict[str, Callable[[bytes], str]] = {
        ".txt": _extract_txt,
        ".docx": _extract_docx,
        ".doc": _extract_doc,
        ".rtf": _extract_rtf,
        ".pdf": _extract_pdf,
    }
    text = extractors[ext](data)
    text = _cleanup(text)
    if not text.strip():
        raise DocumentExtractError("Не удалось извлечь текст из файла")
    return text


def _cleanup(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _extract_txt(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "cp1251", "cp866", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _extract_docx(data: bytes) -> str:
    try:
        from docx import Document
    except ImportError as e:
        raise DocumentExtractError("Нужен пакет python-docx") from e

    doc = Document(io.BytesIO(data))
    parts: list[str] = []
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t:
            parts.append(t)
    for table in doc.tables:
        for row in table.rows:
            cells = [(c.text or "").strip() for c in row.cells]
            cells = [c for c in cells if c]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _extract_rtf(data: bytes) -> str:
    try:
        from striprtf.striprtf import rtf_to_text
    except ImportError as e:
        raise DocumentExtractError("Нужен пакет striprtf") from e

    raw = None
    for enc in ("utf-8", "cp1251", "latin-1"):
        try:
            raw = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if raw is None:
        raw = data.decode("latin-1", errors="replace")
    return rtf_to_text(raw)


def _extract_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as e:
        raise DocumentExtractError("Нужен пакет pypdf") from e

    reader = PdfReader(io.BytesIO(data))
    parts: list[str] = []
    for page in reader.pages:
        t = page.extract_text() or ""
        t = t.strip()
        if t:
            parts.append(t)
    return "\n\n".join(parts)


def _extract_doc(data: bytes) -> str:
    """Legacy .doc: Word COM on Windows, then antiword/catdoc if available."""
    errors: list[str] = []

    try:
        return _extract_doc_win32(data)
    except Exception as e:
        errors.append(f"Word COM: {e}")

    for tool in ("antiword", "catdoc"):
        try:
            return _extract_doc_cli(data, tool)
        except Exception as e:
            errors.append(f"{tool}: {e}")

    hint = (
        "Старый формат .doc не удалось прочитать автоматически. "
        "Сохраните документ как .docx / .txt / .pdf и загрузите снова."
    )
    if errors:
        hint += " (" + "; ".join(errors[:2]) + ")"
    raise DocumentExtractError(hint)


def _extract_doc_win32(data: bytes) -> str:
    import pythoncom  # type: ignore
    import win32com.client  # type: ignore

    pythoncom.CoInitialize()
    word = None
    doc = None
    tmp_path = None
    try:
        fd, tmp_path = tempfile.mkstemp(suffix=".doc")
        os.close(fd)
        Path(tmp_path).write_bytes(data)

        word = win32com.client.Dispatch("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        doc = word.Documents.Open(str(Path(tmp_path).resolve()), ReadOnly=True)
        text = doc.Content.Text or ""
        return text.replace("\r", "\n")
    finally:
        try:
            if doc is not None:
                doc.Close(False)
        except Exception:
            pass
        try:
            if word is not None:
                word.Quit()
        except Exception:
            pass
        try:
            pythoncom.CoUninitialize()
        except Exception:
            pass
        if tmp_path:
            Path(tmp_path).unlink(missing_ok=True)


def _extract_doc_cli(data: bytes, tool: str) -> str:
    with tempfile.NamedTemporaryFile(suffix=".doc", delete=False) as tmp:
        tmp.write(data)
        tmp_path = tmp.name
    try:
        proc = subprocess.run(
            [tool, tmp_path],
            capture_output=True,
            check=False,
            timeout=60,
        )
        if proc.returncode != 0:
            err = (proc.stderr or proc.stdout or b"").decode("utf-8", errors="replace")
            raise RuntimeError(err.strip() or f"{tool} failed")
        out = proc.stdout.decode("utf-8", errors="replace")
        if not out.strip():
            raise RuntimeError("пустой вывод")
        return out
    finally:
        Path(tmp_path).unlink(missing_ok=True)
